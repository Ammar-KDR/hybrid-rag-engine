from pathlib import Path
import re
import unicodedata

from bs4 import BeautifulSoup
from pypdf import PdfReader

from .extraction_audit import audit_extracted_page
from .model import Document
from .table_serializer import (
    extract_and_serialize_pymupdf_tables,
    serialize_table_to_records,
)


SUPPORTED_EXTENSIONS = {".txt", ".md", ".html", ".htm", ".pdf", ".docx"}


MUNICIPAL_FONT_REPAIRS = [
    (re.compile(r"الراض\s*ي?اء"), "الكهرباء"),
    (re.compile(r"\bة\s+الكهرباء"), "شركة الكهرباء"),
    (re.compile(r"\bشيف\b"), "كشف"),
    (re.compile(r"\bاسيباة\b"), "استهلاك"),
    (re.compile(r"\bاالدة\b"), "الإغلاق"),
    (re.compile(r"ادة\s+المنش"), "إغلاق المنشأة"),
    (re.compile(r"\bادلق\b"), "أغلق"),
    (re.compile(r"\bمب\s*اقد\b"), "متعاقد"),
    (re.compile(r"\bال\s*قيد\b"), "العقد"),
    (re.compile(r"\bت\s*ااقد\b"), "تعاقد"),
    (re.compile(r"ر\s*ي(?:ة)?\s+الماين"), "رخصة المهن"),
    (re.compile(r"بيا\s+ت\s*طنية"), "كتاب تغطية"),
    (re.compile(r"الوصيوالت\s+المالنية"), "الوصلات المالية"),
    (re.compile(r"شيف\s+أيا"), "كشف حساب"),
    (re.compile(r"ال\s*ي\s*ة\s+الناقلية?"), "الشركة الناقلة"),
    (re.compile(r"\bأعما\b"), "الأعمال"),
    (re.compile(r"مؤسأات\s+نراية"), "مؤسسات فردية"),
    (re.compile(r"م\s*اتب\s+ماننة"), "مكاتب مهنية"),
    (re.compile(r"المناطق\s+ال\s*ر"), "المناطق الحرة"),
    (re.compile(r"جم\s*نات\s+ت\s*اوننة"), "جمعيات تعاونية"),
]


def normalize_text(text: str) -> str:
    """
    Normalize line endings and Unicode canonical compatibility (NFKC).
    Converts Arabic presentation forms / ligatures to standard Unicode characters,
    fixes reversed visual Arabic from government PDF exporters, strips excessive tatweels,
    and repairs known PDF font encoding corruptions.
    """
    import unicodedata
    import re
    from rag.nlp.arabic import fix_visual_arabic

    text = (
        unicodedata.normalize("NFKC", text)
        .replace("\r\n", "\n")
        .replace("\r", "\n")
    )
    text = fix_visual_arabic(text)
    # Strip tatweel and collapse excessive repeated letters (>=3 repetitions)
    text = re.sub(r"ـ+", "", text)
    text = re.sub(r"([\u0600-\u06FF])\1{2,}", r"\1", text)
    for pattern, repl in MUNICIPAL_FONT_REPAIRS:
        text = pattern.sub(repl, text)
    return text.strip()


def load_text_document(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8")

    return [
        Document(
            text=normalize_text(text),
            source=str(path),
            file_type=path.suffix.lower().lstrip("."),
        )
    ]


def load_html_document(path: Path) -> list[Document]:
    html = path.read_text(encoding="utf-8")

    soup = BeautifulSoup(html, "lxml")

    # Decompose non-content elements and website boilerplate
    for element in soup(["script", "style", "noscript", "meta", "link", "input", "nav", "header", "footer"]):
        element.decompose()
    for sel in [".mean-container", ".mobile-menu", ".top-header", ".main-header", ".footer", ".breadcrumb"]:
        for el in soup.select(sel):
            el.decompose()

    text = soup.get_text(separator="\n", strip=True)

    return [
        Document(
            text=normalize_text(text),
            source=str(path),
            file_type="html",
        )
    ]


def load_pdf_document(path: Path) -> list[Document]:
    # Try pymupdf (fitz) first for fast, robust extraction, table detection, and layout preservation
    try:
        import pymupdf

        doc = pymupdf.open(str(path))
        documents = []
        page_count = len(doc)
        for page_index in range(page_count):
            page = doc[page_index]
            page_num = page_index + 1

            # 1. Native structured table extraction
            table_text, table_count = extract_and_serialize_pymupdf_tables(page, page_num)

            # 2. Page raw text extraction
            raw_text = page.get_text() or ""
            raw_text = normalize_text(raw_text)

            # Interleave/append serialized tables so fees and rules are explicitly bound
            if table_text:
                if raw_text:
                    page_text = f"{raw_text}\n\n{table_text}"
                else:
                    page_text = table_text
            else:
                page_text = raw_text

            # Check for images/scans
            has_images = bool(page.get_images())

            # 3. Page-level extraction audit
            audit = audit_extracted_page(
                filename=path.name,
                page_num=page_num,
                text=page_text,
                method="digital_pdf",
                tables_count=table_count,
                has_images=has_images,
            )

            # Do not silently drop scanned or blank pages
            if not page_text.strip():
                if audit.status == "needs_ocr":
                    page_text = f"[تنبيه: صفحة {page_num} تحتوي على مسح ضوئي أو صورة بدون نص رقمي وتتطلب تدقيقاً يدوياً أو تفعيل OCR]"
                else:
                    page_text = f"[صفحة {page_num} فارغة]"

            documents.append(
                Document(
                    text=page_text,
                    source=str(path),
                    file_type="pdf",
                    metadata={
                        "page_number": page_num,
                        "page_count": page_count,
                        "extraction_method": audit.method if audit.status != "needs_ocr" else "needs_ocr",
                        "extraction_status": audit.status,
                        "audit": audit.to_dict(),
                        "warnings": audit.warnings,
                    },
                )
            )
        doc.close()
        if documents:
            return documents
    except Exception:
        pass

    # Fallback to pypdf
    reader = PdfReader(path)
    documents = []
    page_count = len(reader.pages)

    for page_index, page in enumerate(reader.pages):
        page_num = page_index + 1
        text = page.extract_text() or ""
        text = normalize_text(text)

        audit = audit_extracted_page(
            filename=path.name,
            page_num=page_num,
            text=text,
            method="pypdf",
        )

        if not text.strip():
            text = f"[صفحة {page_num} فارغة]"

        documents.append(
            Document(
                text=text,
                source=str(path),
                file_type="pdf",
                metadata={
                    "page_number": page_num,
                    "page_count": page_count,
                    "extraction_method": "pypdf",
                    "extraction_status": audit.status,
                    "audit": audit.to_dict(),
                    "warnings": audit.warnings,
                },
            )
        )

    return documents


def load_docx_document(path: Path) -> list[Document]:
    import docx

    doc = docx.Document(path)
    elements = []
    table_idx = 0

    # Iterate XML body children in true document order (paragraphs and tables interleaved)
    for child in doc.element.body:
        if child.tag.endswith("p"):
            p = docx.text.paragraph.Paragraph(child, doc)
            p_text = p.text.strip()
            if p_text:
                elements.append(p_text)
        elif child.tag.endswith("tbl"):
            table_idx += 1
            tbl = docx.table.Table(child, doc)
            if tbl.rows:
                headers = [cell.text.strip() for cell in tbl.rows[0].cells]
                data_rows = (
                    [[cell.text.strip() for cell in row.cells] for row in tbl.rows[1:]]
                    if len(tbl.rows) > 1
                    else []
                )

                if not any(headers) and data_rows:
                    headers = [f"عمود {i+1}" for i in range(len(tbl.rows[0].cells))]
                    data_rows = [[cell.text.strip() for cell in tbl.rows[0].cells]]
                elif len(tbl.rows) == 1 and any(headers):
                    data_rows = [headers]
                    headers = [f"عمود {i+1}" for i in range(len(headers))]

                serialized = serialize_table_to_records(headers, data_rows, table_index=table_idx)
                if serialized:
                    elements.append(serialized)

    text = "\n\n".join(elements)
    text = normalize_text(text)

    if not text:
        text = "[مستند Word فارغ]"

    audit = audit_extracted_page(
        filename=path.name,
        page_num=1,
        text=text,
        method="docx",
        tables_count=table_idx,
        has_images=False,
    )

    return [
        Document(
            text=text,
            source=str(path),
            file_type="docx",
            metadata={
                "extraction_method": "docx",
                "extraction_status": audit.status,
                "audit": audit.to_dict(),
                "warnings": audit.warnings,
            },
        )
    ]


def load_document(file_path: str | Path) -> list[Document]:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File does not exist: {path}"
        )

    if not path.is_file():
        raise ValueError(
            f"Expected a file, received: {path}"
        )

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{extension}'. "
            f"Supported types: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    if extension in {".txt", ".md"}:
        return load_text_document(path)

    if extension in {".html", ".htm"}:
        return load_html_document(path)

    if extension == ".pdf":
        return load_pdf_document(path)

    if extension == ".docx":
        return load_docx_document(path)

    # Defensive fallback.
    raise ValueError(
        f"No loader configured for: {extension}"
    )