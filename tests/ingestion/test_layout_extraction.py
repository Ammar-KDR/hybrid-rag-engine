import tempfile
from pathlib import Path
import docx
import pymupdf
import pytest

from rag.ingestion.loader import load_docx_document, load_pdf_document
from rag.ingestion.table_serializer import (
    serialize_table_to_records,
)
from rag.ingestion.extraction_audit import (
    audit_extracted_page,
    calculate_arabic_ratio,
)


def test_table_serializer_record_format():
    headers = ["الخدمة", "الحالة", "الرسوم"]
    rows = [
        ["إصدار بطاقة شخصية", "المرة الأولى", "2 دينار"],
        ["إصدار بدل فاقد", "المرة الثانية", "50 دينار"],
    ]
    result = serialize_table_to_records(headers, rows, table_index=1)
    
    assert "[جدول 1]" in result
    assert "الخدمة: إصدار بطاقة شخصية | الحالة: المرة الأولى | الرسوم: 2 دينار" in result
    assert "الخدمة: إصدار بدل فاقد | الحالة: المرة الثانية | الرسوم: 50 دينار" in result
    assert "| الخدمة | الحالة | الرسوم |" in result


def test_arabic_ratio_calculation():
    pure_arabic = "خدمة إصدار وتجديد جوازات السفر للمواطنين"
    ratio = calculate_arabic_ratio(pure_arabic)
    assert ratio > 0.85

    mixed = "Service Name: بدل فاقد 2026"
    mixed_ratio = calculate_arabic_ratio(mixed)
    assert 0.15 < mixed_ratio < 0.85

    english_only = "Service fee is 50 JOD"
    eng_ratio = calculate_arabic_ratio(english_only)
    assert eng_ratio == 0.0


def test_page_audit_status():
    # 1. OK status
    rec_ok = audit_extracted_page(
        filename="test.pdf",
        page_num=1,
        text="خدمة إصدار وتجديد جواز السفر الأردني مع كافة الرسوم والمتطلبات الرسمية",
        method="digital_pdf",
        tables_count=1,
    )
    assert rec_ok.status == "ok"
    assert rec_ok.tables == 1
    assert rec_ok.chars > 30

    # 2. Scanned / image only page
    rec_scanned = audit_extracted_page(
        filename="test.pdf",
        page_num=2,
        text="",
        method="digital_pdf",
        has_images=True,
    )
    assert rec_scanned.status == "needs_ocr"
    assert any("OCR" in w for w in rec_scanned.warnings)

    # 3. Empty page
    rec_empty = audit_extracted_page(
        filename="test.pdf",
        page_num=3,
        text="",
        method="digital_pdf",
        has_images=False,
    )
    assert rec_empty.status == "empty"


def test_docx_interleaved_document_order():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "interleaved.docx"
        doc = docx.Document()
        
        # Paragraph 1
        doc.add_paragraph("الفقرة الأولى: مقدمة عن الخدمة الحكومية")
        
        # Interleaved Table
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "الخدمة"
        table.rows[0].cells[1].text = "الرسوم"
        table.rows[1].cells[0].text = "تجديد الجواز"
        table.rows[1].cells[1].text = "25 دينار"
        
        # Paragraph 2
        doc.add_paragraph("الفقرة الثانية: الملاحظات والاستثناءات القانونية")
        
        doc.save(str(file_path))
        
        documents = load_docx_document(file_path)
        assert len(documents) == 1
        text = documents[0].text
        
        # Verify sequence: Para 1 appears BEFORE Table, Table appears BEFORE Para 2
        pos_para1 = text.find("الفقرة الأولى")
        pos_table = text.find("[جدول 1]")
        pos_para2 = text.find("الفقرة الثانية")
        
        assert pos_para1 != -1
        assert pos_table != -1
        assert pos_para2 != -1
        assert pos_para1 < pos_table < pos_para2
        
        # Verify row key-value binding
        assert "الخدمة: تجديد الجواز | الرسوم: 25 دينار" in text
        assert documents[0].metadata["extraction_method"] == "docx"
        assert documents[0].metadata["extraction_status"] == "ok"


def test_pdf_extraction_and_page_preservation():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "sample_service.pdf"
        doc = pymupdf.open()

        # Page 1: Text content
        p1 = doc.new_page()
        p1.insert_text((50, 50), "Service: National Identity Card Issuance", fontname="helv")

        # Page 2: Blank page (must NOT be silently dropped)
        p2 = doc.new_page()

        doc.save(str(file_path))
        doc.close()

        documents = load_pdf_document(file_path)

        # Assert both pages are retained (page count parity: 2 pages)
        assert len(documents) == 2
        assert documents[0].metadata["page_number"] == 1
        assert "National Identity Card" in documents[0].text
        assert documents[0].metadata["extraction_method"] == "digital_pdf"
        assert documents[0].metadata["extraction_status"] == "ok"

        assert documents[1].metadata["page_number"] == 2
        assert "[صفحة 2 فارغة]" in documents[1].text
        assert documents[1].metadata["extraction_status"] == "empty"
